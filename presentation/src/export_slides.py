"""Export the approved composition as movable video panels and editable headings."""

from concurrent.futures import ThreadPoolExecutor
import json
import hashlib
import subprocess
import cv2
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation
from pptx.util import Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from reportlab.pdfgen.canvas import Canvas
from reportlab.lib.utils import ImageReader
from paths import OUTPUT, FONTS
from captions import CUES

W, H = 1280, 720
BG = "#f4f6f7"
INK = "#21333e"
TEAL = "#167b76"
MUTED = "#566975"
# Coordinates are the original 1280 x 720 composition, without scaling or cropping.
SLIDES = [
    dict(
        name="opening",
        duration=13,
        poster=12.6,
        top=119,
        panels=[
            ("Pitch", (94, 134, 468, 196)),
            ("Observe", (92, 358, 470, 70)),
            ("Recover", (92, 450, 470, 100)),
            ("Plan", (92, 570, 470, 66)),
            ("Plasticine", (672, 120, 166, 156)),
            ("Play-Doh", (844, 120, 164, 156)),
            ("Butter slime", (1014, 120, 166, 156)),
            ("Press A", (672, 288, 166, 124)),
            ("Press B", (844, 288, 164, 124)),
            ("Press C", (1014, 288, 166, 124)),
            ("Glycerin", (672, 424, 166, 224)),
            ("90% glycerin", (844, 424, 164, 224)),
            ("Water", (1014, 424, 166, 224)),
        ],
    ),
    dict(
        name="observe",
        duration=32,
        poster=31.6,
        top=75,
        panels=[
            ("Recorded deformation", (44, 140, 270, 170)),
            ("Cameras", (66, 310, 232, 118)),
            ("Simulation deformation", (44, 428, 270, 160)),
            ("RGBD surface", (382, 150, 174, 194)),
            ("Volume constraint", (582, 150, 170, 194)),
            ("Particle advection", (772, 150, 174, 194)),
            ("Texture tracking", (382, 410, 188, 204)),
            ("Surface triangulation", (584, 410, 164, 204)),
            ("Interior motion", (774, 410, 172, 204)),
            ("Reconstructed motion", (1018, 153, 214, 194)),
            ("Measured force", (1018, 435, 214, 176)),
        ],
    ),
    dict(
        name="balance",
        duration=28,
        poster=27.6,
        top=75,
        panels=[
            ("Motion and contact force", (66, 154, 160, 228)),
            ("Material state", (250, 154, 118, 136)),
            ("Stress basis", (378, 154, 198, 136)),
            ("Stress model", (250, 290, 326, 90)),
            ("Weak balance", (66, 390, 514, 112)),
            ("Particle integration", (66, 516, 214, 132)),
            ("Particle volume", (300, 516, 126, 132)),
            ("Test field", (436, 516, 148, 132)),
            ("Identified law", (640, 124, 230, 156)),
            ("Linear least squares", (640, 310, 230, 338)),
            ("MPM plan", (934, 124, 286, 256)),
            ("Hardware execution", (934, 432, 286, 208)),
        ],
    ),
    dict(
        name="insertion",
        duration=12,
        poster=11.6,
        top=121,
        panels=[
            ("Bending identification", (44, 220, 172, 324)),
            ("Identified material laws", (260, 238, 236, 366)),
            ("A matched", (544, 202, 338, 222)),
            ("A swapped", (898, 202, 338, 222)),
            ("B matched", (544, 431, 338, 222)),
            ("B swapped", (898, 431, 338, 222)),
        ],
    ),
    dict(
        name="golf",
        duration=10,
        poster=9.6,
        top=121,
        panels=[
            ("A matched", (84, 198, 568, 226)),
            ("A swapped", (668, 198, 568, 226)),
            ("B matched", (84, 428, 568, 226)),
            ("B swapped", (668, 428, 568, 226)),
        ],
    ),
    dict(
        name="simshape",
        duration=15,
        poster=14.6,
        top=121,
        panels=[
            ("Pressing identification", (44, 172, 496, 210)),
            ("Identified material laws", (44, 410, 496, 240)),
            ("A matched", (678, 178, 242, 230)),
            ("A swapped", (958, 178, 242, 230)),
            ("B matched", (678, 424, 242, 230)),
            ("B swapped", (958, 424, 242, 230)),
        ],
    ),
    dict(
        name="hardware",
        duration=30,
        poster=29.6,
        top=121,
        panels=[
            ("Play-Doh observed", (44, 182, 224, 130)),
            ("Butter slime observed", (280, 182, 224, 130)),
            ("Plasticine observed", (516, 182, 224, 130)),
            ("Play-Doh predicted", (44, 343, 224, 130)),
            ("Butter slime predicted", (280, 343, 224, 130)),
            ("Plasticine predicted", (516, 343, 224, 130)),
            ("Play-Doh shaped", (44, 505, 224, 152)),
            ("Butter slime shaped", (280, 505, 224, 152)),
            ("Plasticine shaped", (516, 505, 224, 152)),
            ("Identified laws", (812, 180, 424, 204)),
            ("Execute shaping", (812, 446, 424, 212)),
        ],
    ),
    dict(
        name="pouring",
        duration=26,
        poster=25.6,
        top=121,
        panels=[
            ("Recorded pour", (64, 166, 270, 192)),
            ("Identified viscosity", (446, 163, 388, 192)),
            ("MPM replay", (948, 166, 270, 192)),
            ("MPM planning", (48, 452, 258, 210)),
            ("Selected tilt and command transfer", (322, 442, 394, 82)),
            ("Execution cups", (394, 530, 428, 124)),
            ("Measured volume", (900, 435, 340, 226)),
        ],
    ),
    dict(
        name="takeaway",
        duration=12,
        poster=11.6,
        top=75,
        panels=[
            ("One interaction", (42, 80, 584, 116)),
            ("No backpropagation", (648, 80, 590, 116)),
            ("Use in MPM", (42, 200, 584, 118)),
            ("Transfer", (648, 200, 590, 118)),
            ("Benchmark comparison", (42, 330, 1198, 320)),
        ],
    ),
]
TITLES = {
    "observe": ("FORM / 01: ", "From visual observations to 3D motion"),
    "balance": ("FORM / 02: ", "Recover the material law. Plan the robot action."),
    "insertion": ("Results · Simulation / ", "Elastic rod insertion"),
    "golf": ("Results · Simulation / ", "Putting with a flexible club"),
    "simshape": ("Results · Simulation / ", "Plastic shaping"),
    "hardware": ("Results · Hardware / ", "Pressing and shaping"),
    "pouring": ("Results · Hardware / ", "Pouring a target volume"),
    "takeaway": ("FORM / ", "From one interaction to new robot actions"),
}
SUBTITLES = {
    "insertion": "Identify the material law from bending. Use it to insert the rod through the hole by controlling gripper height and tilt.",
    "golf": "Reuse the previously identified material laws. Plan forward-stroke duration and aim angle to stop the ball in the target.",
    "simshape": "Identify stiffness and yield stress from pressing. Plan six pinches to shape a larger block into an X.",
    "hardware": "Identify material laws from pressing. Plan and execute shaping of fresh specimens into an X.",
    "pouring": "Identify from one glycerol pour. Reuse the model to plan the tilt for each target volume.",
}


def px(v):
    return round(v * 9525)


def snapshot(path, t):
    cap = cv2.VideoCapture(str(path))
    cap.set(cv2.CAP_PROP_POS_MSEC, t * 1000)
    ok, frame = cap.read()
    cap.release()
    if not ok:
        raise RuntimeError(f"Cannot read {path} at {t}")
    return Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))


def native_text(slide, text, x, y, size, color=INK, bold=False, width=None):
    font = ImageFont.truetype(str(FONTS / ("Lato-Bold.ttf" if bold else "Lato-Regular.ttf")), size)
    # Match PIL's default anchor (ascender) with an unpadded PowerPoint text box.
    width = width or font.getlength(text) + 12
    box = slide.shapes.add_textbox(px(x), px(y - 0.18 * size), px(width), px(size * 1.45))
    box.name = "Editable: " + text
    tf = box.text_frame
    tf.clear()
    tf.word_wrap = False
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = MSO_ANCHOR.TOP
    p = tf.paragraphs[0]
    p._p.get_or_add_pPr().set("marL", "0")
    p._p.get_or_add_pPr().set("indent", "0")
    p.space_before = p.space_after = Pt(0)
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = text
    r.font.name = "Lato"
    r.font.size = Pt(size * 0.75)
    r.font.bold = bold
    r.font.color.rgb = RGBColor.from_string(color.lstrip("#"))
    return font.getlength(text)


def editable_header(slide, spec, poster, folder):
    image = poster.crop((0, 0, W, spec["top"]))
    d = ImageDraw.Draw(image)
    if spec["name"] == "opening":
        d.rectangle((30, 10, 1250, 103), fill=BG)
        header = [
            ("FORM", 44, 21, 54, TEAL, True),
            (
                "Robot Manipulation through Direct Material Law Identification",
                227,
                36,
                36,
                INK,
                True,
            ),
            ("From Observed Response to Material laws", 44, 77, 20, TEAL, False),
        ]
    else:
        d.rectangle((30, 10, 1250, 61), fill=BG)
        prefix, title = TITLES[spec["name"]]
        f = ImageFont.truetype(str(FONTS / "Lato-Bold.ttf"), 34)
        header = [
            (prefix, 44, 22, 34, TEAL, True),
            (
                title,
                44 + f.getlength(prefix) + (2 if spec["name"] == "takeaway" else 0),
                22,
                34,
                INK,
                True,
            ),
        ]
        if spec["name"] in SUBTITLES:
            d.rectangle((30, 72, 1250, 111), fill=BG)
            header.append((SUBTITLES[spec["name"]], 44, 79, 23, MUTED, False))
    path = folder / "header.png"
    image.save(path)
    slide.shapes.add_picture(
        str(path), 0, 0, width=px(W), height=px(spec["top"])
    ).name = "Header rule and background"
    for args in header:
        native_text(slide, *args)


def encode(source, destination, filtergraph):
    subprocess.run(
        [
            "ffmpeg",
            "-y",
            "-v",
            "error",
            "-i",
            str(source),
            "-vf",
            filtergraph,
            "-an",
            "-c:v",
            "libx264",
            "-preset",
            "fast",
            "-crf",
            "12",
            "-threads",
            "2",
            "-pix_fmt",
            "yuv420p",
            "-map_metadata",
            "-1",
            "-movflags",
            "+faststart",
            str(destination),
        ],
        check=True,
    )


def prepare_media(spec):
    folder = OUTPUT / "slide_media" / spec["name"]
    folder.mkdir(parents=True, exist_ok=True)
    source = OUTPUT / "sections" / f"{spec['name']}.mp4"
    poster = snapshot(source, spec["poster"])
    stilldir = OUTPUT / "slides"
    stilldir.mkdir(exist_ok=True)
    poster.save(stilldir / f"{SLIDES.index(spec) + 1:02d}_{spec['name']}.png")
    # Residual layout carries synchronized reveals and arrows. Panel rectangles
    # are erased so moving a panel does not expose a second copy underneath.
    top = spec["top"]
    height = H - top
    filters = [
        f"drawbox=x={x}:y={y}:w={w}:h={h}:color=0xf4f6f7:t=fill"
        for _, (x, y, w, h) in spec["panels"]
    ]
    filters.append("drawbox=x=0:y=663:w=1280:h=53:color=0xf4f6f7:t=fill")
    # Pad one pixel for yuv420p where the logical body height is odd.
    filters += [
        f"crop=1280:{height}:0:{top}:exact=1",
        f"pad=1280:{height + height % 2}:0:0:color=0xf4f6f7",
    ]
    items = [
        (
            "Layout and animated connectors",
            (0, top, W, height),
            folder / "layout.mp4",
            ",".join(filters),
        )
    ]
    for i, (name, (x, y, w, h)) in enumerate(spec["panels"]):
        assert x >= 0 and y >= top and x + w <= W and y + h <= 663, (spec["name"], name)
        assert w % 2 == h % 2 == 0, (spec["name"], name)
        items.append(
            (name, (x, y, w, h), folder / f"panel_{i:02}.mp4", f"crop={w}:{h}:{x}:{y}:exact=1")
        )
    items.append(
        (
            "Timed conference captions",
            (0, 664, W, 52),
            folder / "captions.mp4",
            "crop=1280:52:0:664:exact=1",
        )
    )
    jobs = []
    for name, bounds, path, filters in items:
        signature = hashlib.sha256((filters + str(source.stat().st_mtime_ns)).encode()).hexdigest()
        signature_path = path.with_suffix(".signature")
        if (
            not path.exists()
            or not signature_path.exists()
            or signature_path.read_text() != signature
        ):
            jobs.append((source, path, filters))

    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(lambda job: encode(*job), jobs))
    for source, path, filters in jobs:
        path.with_suffix(".signature").write_text(
            hashlib.sha256((filters + str(source.stat().st_mtime_ns)).encode()).hexdigest()
        )
    layers = []
    for name, (x, y, w, h), path, _ in items:
        still = folder / (path.stem + ".png")
        if name.startswith("Layout"):
            pic = poster.crop((0, top, W, H))
            d = ImageDraw.Draw(pic)
            for _, (a, b, c, e) in spec["panels"]:
                d.rectangle((a, b - top, a + c - 1, b + e - top - 1), fill=BG)
            d.rectangle((0, 663 - top, W, 715 - top), fill=BG)
            if height % 2:
                padded = Image.new("RGB", (W, height + 1), BG)
                padded.paste(pic)
                pic = padded
        else:
            pic = poster.crop((x, y, x + w, y + h))
        pic.save(still)
        layers.append(
            (name, (x, y, w, h + (height % 2 if name.startswith("Layout") else 0)), path, still)
        )
    return poster, layers, folder


def export():
    OUTPUT.mkdir(exist_ok=True)
    prs = Presentation()
    prs.slide_width = px(W)
    prs.slide_height = px(H)
    from datetime import datetime

    prs.core_properties.created = prs.core_properties.modified = datetime(2026, 9, 23)
    prs.core_properties.title = "FORM: Video Presentation"
    prs.core_properties.subject = "Reusable slides with Python-rendered video panels"
    prs.core_properties.author = ""
    prs.core_properties.last_modified_by = ""
    prs.core_properties.comments = ""
    pdf = Canvas(str(OUTPUT / "FORM_slides.pdf"), pagesize=(960, 540))
    pdf.setTitle("FORM: Presentation slides")
    pdf.setAuthor("")
    manifest = []
    for spec in SLIDES:
        print("Exporting", spec["name"], flush=True)
        poster, layers, folder = prepare_media(spec)
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        slide.background.fill.solid()
        slide.background.fill.fore_color.rgb = RGBColor.from_string("F4F6F7")
        for name, (x, y, w, h), path, still in layers:
            if spec["name"] == "opening" and name not in (
                "Observe",
                "Recover",
                "Plan",
                "Layout and animated connectors",
                "Timed conference captions",
            ):
                shape = slide.shapes.add_picture(
                    str(still), px(x), px(y), width=px(w), height=px(h)
                )
                shape.name = name
                continue
            shape = slide.shapes.add_movie(
                str(path),
                px(x),
                px(y),
                px(w),
                px(h),
                poster_frame_image=str(still),
                mime_type="video/mp4",
            )
            shape.name = name
        editable_header(slide, spec, poster, folder)
        # Start all panel videos together, rather than requiring individual clicks.
        for cond in slide._element.xpath(".//p:video/p:cMediaNode/p:cTn/p:stCondLst/p:cond"):
            cond.set("delay", "0")
        transition = OxmlElement("p:transition")
        transition.set("advClick", "1")
        transition.set("advTm", str(spec["duration"] * 1000))
        timing = slide._element.find(
            "{http://schemas.openxmlformats.org/presentationml/2006/main}timing"
        )
        slide._element.insert(list(slide._element).index(timing), transition)
        slide.notes_slide.notes_text_frame.text = "\n".join(
            [
                spec["name"].capitalize(),
                f"Duration: {spec['duration']} seconds. Each video panel is independently movable and copyable.",
                "Narration:",
            ]
            + [f"{a:g}–{b:g}s: {t}" for a, b, t in CUES[spec["name"]]]
        )
        pdf.drawImage(ImageReader(poster), 0, 0, 960, 540)
        pdf.showPage()
        manifest.append(
            {
                "slide": spec["name"],
                "duration": spec["duration"],
                "objects": [
                    {"name": n, "bounds": b, "media": str(p.relative_to(OUTPUT))}
                    for n, b, p, _ in layers
                ],
            }
        )
    pdf.save()
    prs.save(OUTPUT / "FORM_slides.pptx")
    (OUTPUT / "slide_objects.json").write_text(json.dumps(manifest, indent=2) + "\n")
    print("Wrote FORM_slides.pptx and FORM_slides.pdf", flush=True)
